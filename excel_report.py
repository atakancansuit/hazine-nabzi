"""Aylık yönetim raporunu Excel dosyası olarak üretir.

Veriyi SQL görünümlerinden alır; hesap yapmaz, yalnızca biçimlendirir. Dosya
reports/ altına ay damgasıyla yazılır: reports/rapor_2026-08.xlsx

Dört sayfa:
  Özet      ana kalemler, göstergeler ve yıl sonu tahmini
  Kalemler  gider detayındaki 23 kalem
  Kurumlar  genel bütçeli idareler
  Aylık     ana kalemlerin ay ay gerçekleşmesi (pivot kurmak isteyen için)

Kullanım:
    python excel_report.py
    python excel_report.py --year 2025
"""
import argparse
import warnings
from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from load import connect

ROOT = Path(__file__).parent
REPORT_DIR = ROOT / "reports"

# Muhasebat'ın kurumsal renkleri; Power BI teması da aynı paleti kullanıyor.
NAVY = "222D56"
GOLD = "C7A069"
LIGHT = "F8F9FA"
RED = "C0392B"
GREEN = "1E7B45"

TITLE_FONT = Font(name="Segoe UI", size=14, bold=True, color=NAVY)
NOTE_FONT = Font(name="Segoe UI", size=9, color="6C757D")
HEAD_FONT = Font(name="Segoe UI", size=10, bold=True, color="FFFFFF")
HEAD_FILL = PatternFill("solid", fgColor=NAVY)
THIN = Side(style="thin", color="DEE2E6")


def sql(connection, query, **params):
    # pandas, SQLAlchemy dışındaki bağlantılar için uyarı veriyor. Bağlantı
    # pymssql ile kuruluyor ve sorunsuz çalışıyor; uyarı bastırılıyor.
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        return pd.read_sql(query.format(**params), connection)


def write_sheet(writer, name, frame, title, note, formats):
    """Bir DataFrame'i biçimlendirilmiş sayfa olarak yazar.

    formats: {sütun adı: excel sayı biçimi}. Sütun genişlikleri içeriğe göre
    ayarlanır, başlık satırı dondurulur, planı aşan satırlar kırmızıya boyanır.
    """
    frame.to_excel(writer, sheet_name=name, startrow=3, index=False)
    sheet = writer.sheets[name]

    sheet["A1"] = title
    sheet["A1"].font = TITLE_FONT
    sheet["A2"] = note
    sheet["A2"].font = NOTE_FONT

    for column, header in enumerate(frame.columns, start=1):
        cell = sheet.cell(row=4, column=column, value=header)
        cell.font, cell.fill = HEAD_FONT, HEAD_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        letter = get_column_letter(column)
        width = max(len(str(header)), *(len(str(v)) for v in frame[header])) + 2
        sheet.column_dimensions[letter].width = min(max(width, 10), 45)
        for row in range(5, len(frame) + 5):
            value = sheet.cell(row=row, column=column)
            value.border = Border(bottom=THIN)
            if header in formats:
                value.number_format = formats[header]
            # Gerçekleşme oranı: planı aşan kırmızı, altında kalan yeşil.
            if header.startswith("Gerçekleşme") and isinstance(value.value, (int, float)):
                value.font = Font(color=RED if value.value > 1 else GREEN, bold=True)

    sheet.freeze_panes = "A5"
    sheet.sheet_view.showGridLines = False


def main(year=None):
    """Raporu üretir. year verilmezse veritabanındaki en güncel yıl kullanılır."""
    with connect() as connection:
        year = year or sql(connection, "SELECT MAX(year) AS y FROM v_annual").y[0]
        months = sql(connection, """
            SELECT MAX(months_reported) AS m FROM v_annual
            WHERE source = 'balance' AND year = {year}""", year=year).m[0]

        ozet = sql(connection, """
            SELECT d.display_name                      AS [Kalem],
                   a.actual / 1e6                      AS [Gerçekleşen (milyar TL)],
                   a.planned / 1e6                     AS [Plan (milyar TL)],
                   a.variance / 1e6                    AS [Sapma (milyar TL)],
                   a.pct_of_plan / 100                 AS [Gerçekleşme],
                   f.forecast_year_end / 1e6           AS [Yıl sonu tahmini (milyar TL)],
                   f.forecast_pct_of_plan / 100        AS [Tahmin / plan]
            FROM v_annual a
            JOIN dim_item d ON d.item_key = a.item_key
            LEFT JOIN v_year_end_forecast f ON f.item_key = a.item_key AND f.year = a.year
            WHERE a.source = 'balance' AND a.year = {year}
            ORDER BY ABS(a.variance) DESC""", year=year)

        kalemler = sql(connection, """
            SELECT d.display_name                 AS [Kalem],
                   a.actual / 1e6                 AS [Gerçekleşen (milyar TL)],
                   a.planned / 1e6                AS [Plan (milyar TL)],
                   a.variance / 1e6               AS [Sapma (milyar TL)],
                   a.pct_of_plan / 100            AS [Gerçekleşme],
                   (m.cumulative_pct_of_plan - m.typical_pct_of_plan) / 100 AS [Kıyas farkı]
            FROM v_annual a
            JOIN dim_item d ON d.item_key = a.item_key
            LEFT JOIN v_monthly m
                   ON m.item_key = a.item_key AND m.year = a.year AND m.month = {months}
            WHERE a.source = 'expense_detail' AND a.year = {year} AND d.is_total = 0
            ORDER BY a.pct_of_plan DESC""", year=year, months=months)

        kurumlar = sql(connection, """
            SELECT d.display_name      AS [Kurum],
                   a.actual / 1e6      AS [Gerçekleşen (milyar TL)],
                   a.planned / 1e6     AS [Plan (milyar TL)],
                   a.variance / 1e6    AS [Sapma (milyar TL)],
                   a.pct_of_plan / 100 AS [Gerçekleşme]
            FROM v_annual a
            JOIN dim_item d ON d.item_key = a.item_key
            WHERE a.source = 'ministries' AND a.year = {year}
            ORDER BY a.actual DESC""", year=year)

        aylik = sql(connection, """
            SELECT d.display_name          AS [Kalem],
                   m.month                 AS [Ay],
                   m.actual / 1e6          AS [Gerçekleşen (milyar TL)],
                   m.cumulative_actual / 1e6 AS [Kümülatif (milyar TL)],
                   m.cumulative_pct_of_plan / 100 AS [Kümülatif / plan]
            FROM v_monthly m
            JOIN dim_item d ON d.item_key = m.item_key
            WHERE m.source = 'balance' AND m.year = {year}
            ORDER BY d.display_name, m.month""", year=year)

    REPORT_DIR.mkdir(exist_ok=True)
    path = REPORT_DIR / f"rapor_{year}-{months:02d}.xlsx"
    stamp = f"{year} yılı, {months} aylık veri · hazırlandı {date.today():%d.%m.%Y}"
    amount = '#,##0'
    percent = '0.0%'

    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        write_sheet(writer, "Özet", ozet,
                    "Bütçe gerçekleşme raporu — ana kalemler",
                    stamp + " · Tahmin sütunları yalnızca yıl bitmemişse dolu",
                    {c: amount for c in ozet.columns if "milyar" in c} |
                    {"Gerçekleşme": percent, "Tahmin / plan": percent})
        write_sheet(writer, "Kalemler", kalemler,
                    "Gider kalemleri",
                    stamp + " · Kıyas farkı: bu yılın temposu ile geçmiş yılların aynı ayındaki tempo arasındaki fark",
                    {c: amount for c in kalemler.columns if "milyar" in c} |
                    {"Gerçekleşme": percent, "Kıyas farkı": percent})
        write_sheet(writer, "Kurumlar", kurumlar,
                    "Genel bütçeli idareler",
                    stamp + " · Harcamaya göre sıralı",
                    {c: amount for c in kurumlar.columns if "milyar" in c} |
                    {"Gerçekleşme": percent})
        write_sheet(writer, "Aylık", aylik,
                    "Ana kalemlerin aylık gerçekleşmesi",
                    stamp + " · Pivot kurmak için ham tablo",
                    {c: amount for c in aylik.columns if "milyar" in c} |
                    {"Kümülatif / plan": percent})

    print(f"{path.name}  ({path.stat().st_size / 1024:.0f} KB)  "
          f"{len(ozet)} + {len(kalemler)} + {len(kurumlar)} + {len(aylik)} satır")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--year", type=int, help="raporlanacak yıl (varsayılan: en güncel yıl)")
    main(parser.parse_args().year)
